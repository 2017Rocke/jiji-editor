from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

root = Path("output/fixtures")
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("吉吉 Office 转换测试", 0)
doc.add_paragraph("Word 中文内容与表格应保留。Free PDF conversion.")
table = doc.add_table(rows=2, cols=2)
for cell, value in zip([table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)], ["项目", "状态", "免费转换", "通过"]):
    cell.text = value
doc.save(root / "sample.docx")
book = Workbook()
sheet = book.active
sheet["A1"] = "吉吉转换测试"
sheet["A2"] = "合计"
sheet["B2"] = "=SUM(10,20)"
book.save(root / "sample.xlsx")
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = "吉吉转换测试"
slide.placeholders[1].text = "PowerPoint 转 PDF · 中文内容"
slides.save(root / "sample.pptx")
print("Created synthetic Office fixtures.")
